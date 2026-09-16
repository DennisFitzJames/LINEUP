from __future__ import annotations

import csv
import io
import json
import re
from collections import defaultdict
from datetime import datetime, date, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[2]
CAP_FILE = PROJECT_ROOT / "data" / "processed" / "capacities_clean.csv"
PROD_FILE = PROJECT_ROOT / "data" / "processed" / "production_list_clean.csv"
CONFIG_FILE = PROJECT_ROOT / "config" / "production_plans.json"

DATE_TIME_FORMATS = ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S")
DATE_FORMATS = ("%d.%m.%Y", "%Y-%m-%d")


def clean(value: Any) -> str:
    return str(value or "").replace("\xa0", " ").strip()


def number(value: Any) -> float:
    try:
        return float(clean(value).replace(",", ""))
    except (TypeError, ValueError):
        return 0.0


def parse_datetime(value: Any) -> datetime | None:
    text = clean(value)
    if not text:
        return None
    for fmt in DATE_TIME_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def fmt_dt(value: datetime | None) -> str:
    return value.strftime("%d/%m/%Y %H:%M") if value else ""


def fmt_date(value: datetime | None) -> str:
    return value.strftime("%d/%m/%Y") if value else ""


def slugify(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def status_tokens(value: Any) -> list[str]:
    return re.sub(r"[,;/|-]+", " ", clean(value).upper()).split()


def is_released(system_status: str) -> bool:
    return "REL" in status_tokens(system_status)


def is_teco(system_status: str) -> bool:
    return "TECO" in status_tokens(system_status)


def op_sort(operation: str) -> tuple[int, str]:
    text = clean(operation)
    try:
        return int(float(text)), text
    except ValueError:
        return 999999999, text


def _read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
        return list(csv.DictReader(handle))


def load_config() -> dict[str, Any]:
    with CONFIG_FILE.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def _source_signature() -> tuple[float, int, float, int, float]:
    return (
        CAP_FILE.stat().st_mtime if CAP_FILE.exists() else 0,
        CAP_FILE.stat().st_size if CAP_FILE.exists() else 0,
        PROD_FILE.stat().st_mtime if PROD_FILE.exists() else 0,
        PROD_FILE.stat().st_size if PROD_FILE.exists() else 0,
        CONFIG_FILE.stat().st_mtime if CONFIG_FILE.exists() else 0,
    )


@lru_cache(maxsize=4)
def _build_model_cached(signature: tuple[float, int, float, int, float]) -> dict[str, Any]:
    config = load_config()
    excluded = {clean(x).upper() for x in config.get("excluded_work_centres", [])}
    prod_rows = _read_csv(PROD_FILE)
    cap_rows = _read_csv(CAP_FILE)

    prod_lookup: dict[str, dict[str, str]] = {}
    for row in prod_rows:
        order = clean(row.get("order_number"))
        if order and order not in prod_lookup:
            prod_lookup[order] = row

    all_by_order: dict[str, list[dict[str, Any]]] = defaultdict(list)
    by_wc: dict[str, list[dict[str, Any]]] = defaultdict(list)

    for index, row in enumerate(cap_rows, start=2):
        if clean(row.get("capacity_category")) not in {"2", "002"}:
            continue
        order = clean(row.get("order_number"))
        wc = clean(row.get("work_centre")).upper()
        if not order or not wc:
            continue
        start = parse_datetime(row.get("earliest_start"))
        finish = parse_datetime(row.get("latest_finish"))
        actual_start = parse_datetime(row.get("actual_start"))
        actual_finish = parse_datetime(row.get("actual_finish"))
        prod = prod_lookup.get(order, {})
        rem_proc = number(row.get("remaining_process_hours"))
        rem_setup = number(row.get("remaining_setup_hours"))
        rem_teardown = number(row.get("remaining_teardown_hours"))
        target_proc = number(row.get("target_process_hours"))
        target_setup = number(row.get("target_setup_hours"))
        target_teardown = number(row.get("target_teardown_hours"))
        system_status = clean(prod.get("system_status"))
        completed = abs(rem_proc) < 1e-9 and abs(rem_setup) < 1e-9 and abs(rem_teardown) < 1e-9
        item = {
            "source_row": index,
            "order_number": order,
            "operation": clean(row.get("operation")),
            "work_centre": wc,
            "scheduled_start": start,
            "scheduled_finish": finish,
            "actual_start": actual_start,
            "actual_finish": actual_finish,
            "operation_description": clean(row.get("operation_description")),
            "material": clean(prod.get("material")),
            "material_description": clean(prod.get("material_description")),
            "target_quantity": number(prod.get("target_quantity")),
            "confirmed_quantity": number(prod.get("confirmed_quantity")),
            "customer": clean(prod.get("customer")),
            "priority": clean(prod.get("priority")),
            "system_status": system_status,
            "user_status": clean(prod.get("user_status")),
            "target_process_hours": target_proc,
            "target_setup_hours": target_setup,
            "target_teardown_hours": target_teardown,
            "total_hours": target_proc + target_teardown,
            "setup_hours": target_setup,
            "remaining_hours": rem_proc + rem_setup + rem_teardown,
            "remaining_process_hours": rem_proc,
            "remaining_setup_hours": rem_setup,
            "remaining_teardown_hours": rem_teardown,
            "completed": completed,
            "released": is_released(system_status),
            "teco": is_teco(system_status),
        }
        all_by_order[order].append(item)
        if wc not in excluded:
            by_wc[wc].append(item)

    # VBA sequence is operation first, then date, then source row.
    next_map: dict[int, dict[str, Any]] = {}
    for operations in all_by_order.values():
        operations.sort(key=lambda x: (op_sort(x["operation"]), x["scheduled_start"] or datetime.max, x["source_row"]))
        for current, nxt in zip(operations, operations[1:]):
            next_map[current["source_row"]] = nxt

    plans: dict[str, dict[str, Any]] = {}
    today = date.today()
    horizon = today + timedelta(days=21)
    for wc, rows in by_wc.items():
        rows.sort(key=lambda x: (x["scheduled_start"] or datetime.max, op_sort(x["operation"]), x["order_number"]))
        public_rows = []
        for row in rows:
            nxt = next_map.get(row["source_row"])
            next_wc = nxt["work_centre"] if nxt else ""
            next_available = bool(next_wc and next_wc not in excluded and next_wc in by_wc)
            public = {k: v for k, v in row.items() if k not in {"scheduled_start", "scheduled_finish", "actual_start", "actual_finish"}}
            public.update({
                "scheduled_start_display": fmt_dt(row["scheduled_start"]),
                "scheduled_finish_display": fmt_dt(row["scheduled_finish"]),
                "actual_start_display": fmt_dt(row["actual_start"]),
                "actual_finish_display": fmt_dt(row["actual_finish"]),
                "scheduled_date_key": row["scheduled_start"].strftime("%Y-%m-%d") if row["scheduled_start"] else "unscheduled",
                "scheduled_date_display": row["scheduled_start"].strftime("%A %d/%m/%Y").upper() if row["scheduled_start"] else "UNSCHEDULED",
                "next_work_centre": next_wc,
                "next_available": next_available,
                "next_anchor": f"operation-{nxt['order_number']}-{nxt['operation']}-{slugify(next_wc)}" if nxt else "",
                "row_anchor": f"operation-{row['order_number']}-{row['operation']}-{slugify(wc)}",
                "row_class": "teco" if row["teco"] else "unreleased" if not row["released"] else "completed" if row["completed"] else "active",
            })
            public_rows.append(public)
        active_count = sum(not r["completed"] and not r["teco"] for r in rows)
        completed_count = sum(r["completed"] for r in rows)
        unreleased_count = sum(not r["released"] and not r["teco"] for r in rows)
        remaining_hours = sum(r["remaining_hours"] for r in rows if not r["completed"] and not r["teco"])
        near_term_unreleased = sum(
            (not r["released"] and not r["teco"] and r["scheduled_start"] is not None and today <= r["scheduled_start"].date() <= horizon)
            for r in rows
        )
        plans[wc] = {
            "id": slugify(wc), "work_centre": wc, "rows": public_rows,
            "active_count": active_count, "completed_count": completed_count,
            "unreleased_count": unreleased_count, "remaining_hours": round(remaining_hours, 3),
            "near_term_unreleased": near_term_unreleased,
        }

    work_centres = sorted(plans)
    generated = datetime.now()
    totals = {
        "work_centres": len(work_centres),
        "active_operations": sum(p["active_count"] for p in plans.values()),
        "completed_operations": sum(p["completed_count"] for p in plans.values()),
        "unreleased_operations": sum(p["unreleased_count"] for p in plans.values()),
        "remaining_hours": round(sum(p["remaining_hours"] for p in plans.values()), 3),
    }
    work_centre_areas: dict[str, set[str]] = defaultdict(set)
    for area_name, groups in config.get("export_areas", {}).items():
        for source_work_centres in groups.values():
            for source_wc in source_work_centres:
                work_centre_areas[clean(source_wc).upper()].add(area_name)

    cards = []
    for plan in plans.values():
        card = {k: v for k, v in plan.items() if k != "rows"}
        card["areas"] = sorted(work_centre_areas.get(plan["work_centre"], set()))
        cards.append(card)
    cards.sort(key=lambda p: (-p["active_count"], p["work_centre"]))
    return {
        "generated_at": generated.isoformat(timespec="seconds"),
        "generated_at_display": generated.strftime("%d/%m/%Y %H:%M:%S"),
        "refresh_seconds": int(config.get("refresh_seconds", 60)),
        "plans": plans, "work_centres": work_centres, "cards": cards, "totals": totals,
        "config": config,
    }


def get_model() -> dict[str, Any]:
    return _build_model_cached(_source_signature())


def get_plan(work_centre: str) -> dict[str, Any] | None:
    return get_model()["plans"].get(clean(work_centre).upper())


def export_area_workbook(area_name: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter

    model = get_model()
    export_areas = model["config"].get("export_areas", {})
    area = export_areas.get(area_name)
    if not area:
        raise KeyError(area_name)

    prod_lookup = {clean(r.get("order_number")): r for r in _read_csv(PROD_FILE) if clean(r.get("order_number"))}
    cap_rows = _read_csv(CAP_FILE)
    by_source_wc: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in cap_rows:
        if clean(row.get("capacity_category")) in {"2", "002"}:
            by_source_wc[clean(row.get("work_centre")).upper()].append(row)

    wb = Workbook()
    wb.remove(wb.active)
    headers = ["Scheduled Start", "Scheduled Finish", "Actual Start", "Work Centre", "Order Number", "Material Number", "Material Description", "Order Quantity", "Confirmed Quantity", "Customer Name", "Priority", "System Status", "User Status", "Hours Required", "Setup Time"]
    thin = Side(style="thin", color="D9D9D9")

    for sheet_name, source_wcs in area.items():
        grouped: dict[tuple[str, str], list[Any]] = {}
        order_keys: list[tuple[str, str]] = []
        for source_wc in source_wcs:
            for cap in by_source_wc.get(source_wc.upper(), []):
                order = clean(cap.get("order_number"))
                rem_proc = number(cap.get("remaining_process_hours"))
                rem_setup = number(cap.get("remaining_setup_hours"))
                rem_td = number(cap.get("remaining_teardown_hours"))
                if not order or rem_proc + rem_setup + rem_td <= 1e-9:
                    continue
                key = (source_wc.upper(), order)
                prod = prod_lookup.get(order, {})
                if key not in grouped:
                    grouped[key] = [
                        parse_datetime(prod.get("scheduled_start")), parse_datetime(prod.get("scheduled_finish")), parse_datetime(prod.get("actual_start")),
                        source_wc.upper(), order, clean(prod.get("material")), clean(prod.get("material_description")), number(prod.get("target_quantity")),
                        number(prod.get("confirmed_quantity")), clean(prod.get("customer")), clean(prod.get("priority")), clean(prod.get("system_status")),
                        clean(prod.get("user_status")), 0.0, 0.0,
                    ]
                    order_keys.append(key)
                grouped[key][13] += rem_proc + rem_td
                grouped[key][14] += rem_setup

        ws = wb.create_sheet(title=sheet_name[:31])
        ws.append(headers)
        for key in sorted(order_keys, key=lambda k: (grouped[k][0] or datetime.max, k[0], k[1])):
            ws.append(grouped[key])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="5F5BD6")
            cell.alignment = Alignment(horizontal="center")
        for row in ws.iter_rows():
            for cell in row:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.column_dimensions["E"].number_format = "@"
        for col in ("A", "B", "C"):
            for cell in ws[col][1:]:
                cell.number_format = "DD/MM/YYYY"
        widths = [16,16,16,15,17,18,36,14,18,28,11,28,24,15,13]
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
